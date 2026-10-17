# /// script
# requires-python = ">=3.11"
# dependencies = ["numpy>=1.26", "matplotlib>=3.8", "python-pptx>=1.0"]
# ///
"""Cut rendered videos into clips at the pauses in pauses.toml, and optionally build a deck that plays them.

    uv run animations/cut.py                # every video in pauses.toml: clips + pause sheets
    uv run animations/cut.py intro outro    # only these videos
    uv run animations/cut.py --sheet        # only the pause sheets (quick check of where the cuts fall)
    uv run animations/cut.py --gif          # also each clip as a GIF that plays once and holds its last frame
    uv run animations/cut.py --pptx         # also videos/clips/deck.pptx: one slide per video, one click per clip
    uv run animations/cut.py --pptx --cemep # the deck with the CEMEP title: blue bar, white title, CEMEP symbol
    uv run animations/cut.py --pptx --bar '#1f4e79' --logo animations/data/logo.png   # another bar color, a logo

Renders nothing new and writes only under videos/clips/:
    <video>/01.mp4, 02.mp4, ...   the last frame of each clip is the first frame of the next (seamless)
    <video>_pauses.png            the frame where each clip stops, with its instant
    deck.pptx                     one 16:9 slide per video cut in this run, in pauses.toml order: the video's
                                  `title` from pauses.toml on top (plain, or in a bar with --cemep or --bar) and
                                  under it the video's clips, stacked. Clip 1 plays on the first click; each next
                                  click shows the next clip and plays it at once ("with previous"), so the video
                                  goes on from where it stopped. The notes list each click with its pause comment.
                                  Checked after saving.
"""
import argparse
import itertools
import re
import shutil
import subprocess
import tempfile
import tomllib
from pathlib import Path

from anim import INK, OUT, PAUSES_FILE, sheet, to_gif

DEST = OUT / "clips"
# the deck, in EMU: a 16:9 slide; the title's space on top, 18% of its height; the title's left edge; the logo's
# height and margin. The CEMEP title: this bar color, a white title, the CEMEP symbol (beside this script) at right.
SLIDE_W, SLIDE_H, BAR_H = 12192000, 6858000, 1239520
TITLE_X, LOGO_H, LOGO_MARGIN = 806991, 644550, 359252
BAR = "232751"
CEMEP_LOGO = Path(__file__).with_name("cemep_logo.png")


def run(*cmd):
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", *cmd], check=True)


def probe(path):
    """(frame count, fps, width, height) of a video."""
    out = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries",
                          "stream=nb_frames,r_frame_rate,width,height", "-of", "default=noprint_wrappers=1",
                          str(path)], check=True, capture_output=True, text=True).stdout
    info = dict(line.split("=", 1) for line in out.split())
    num, den = info["r_frame_rate"].split("/")
    return int(info["nb_frames"]), int(num) / int(den), int(info["width"]), int(info["height"])


def bounds(pauses, n, fps, name):
    """Cut frames [0, pause1, ..., last]: clip k runs from bounds[k-1] to bounds[k], both included."""
    frames = [int(round(t * fps)) for t in pauses]
    last = n - 1
    for t, f, prev in zip(pauses, frames, [0] + frames):
        if not 0 < f < last:
            raise SystemExit(f"{name}: pause {t} s is outside the video (0 to {last / fps:.2f} s)")
        if f - prev < 2:
            raise SystemExit(f"{name}: pause {t} s is out of order or glued to the previous one")
    if frames and last - frames[-1] < 2:
        raise SystemExit(f"{name}: the last pause is glued to the end of the video")
    return [0, *frames, last]


def cut(src, a, b, dst):
    run("-i", str(src), "-vf", f"trim=start_frame={a}:end_frame={b + 1},setpts=PTS-STARTPTS", "-an",
        "-c:v", "libx264", "-crf", "18", "-preset", "slow", "-pix_fmt", "yuv420p", "-movflags", "+faststart",
        "-colorspace", "bt709", "-color_primaries", "bt709", "-color_trc", "bt709", str(dst))


def frame_png(src, f, dst):
    run("-i", str(src), "-vf", f"select=eq(n\\,{f})", "-frames:v", "1", str(dst))


PAUSE_ROW = re.compile(r"^\s*[\d.]+\s*,?\s*(?:#\s*(.*?))?\s*$")  # a pause, with or without its comment
COMMENT_ROW = re.compile(r"^\s*#\s*(.*?)\s*$")
CLIP_NUMBER = re.compile(r"^(?:\d  |\d{2,} )")  # what anim.py writes in front of a comment: f"{k:<2} "


def clip_comments(text, name):
    """The comment of each clip of video `name`, in order, from the text of pauses.toml (tomllib drops comments):
    the one on each pause line ('' if none), then the first comment line after the last pause, for the last clip.
    Comment lines between pauses (a deleted number, a pause commented out, a wrapped comment) are skipped, so note k
    stays with clip k. The clip number anim.py writes in front of a comment is left out. [] if there is no section."""
    head = re.compile(r'^\s*\[\s*"?' + re.escape(name) + r'"?\s*\]')
    out, last, where = [], None, "before"
    for line in text.splitlines():
        if where == "before":
            where = "section" if head.match(line) else "before"
        elif line.lstrip().startswith("["):  # the next section
            break
        elif where == "section":
            if re.match(r"\s*pauses\s*=\s*\[", line):
                where = "list" if "]" not in line else "inline"  # an inline list has no comment per pause
        elif where == "inline" or line.strip().startswith("]"):
            break
        elif m := PAUSE_ROW.match(line):
            out.append(m.group(1) or "")
            last = None
        elif (m := COMMENT_ROW.match(line)) and last is None:
            last = m.group(1)
    if where == "before":
        return []
    return [CLIP_NUMBER.sub("", c) for c in out + [last or ""]]


def video_box(vw, vh):
    """(x, y, w, h) in EMU: the largest vw:vh rectangle under the title bar, centered, touching the bar."""
    s = min(SLIDE_W / vw, (SLIDE_H - BAR_H) / vh)
    w, h = round(vw * s), round(vh * s)
    return (SLIDE_W - w) // 2, BAR_H, w, h


def bar_color(s):
    """'#232751' or '232751' -> '232751' (the type of --bar)."""
    if not re.fullmatch(r"#?[0-9A-Fa-f]{6}", s):
        raise argparse.ArgumentTypeError(f"{s!r} is not a color like '#232751'")
    return s.lstrip("#").upper()


def refuse_if_open(deck):
    """PowerPoint keeps ~$<name> beside a deck it has open: writing the deck then would lose the edits made there."""
    if (deck.parent / f"~${deck.name}").exists():
        raise SystemExit(f"{deck.name} is open in PowerPoint: close it (save it under another name to keep any "
                         "edits) and run again")


# A slide's animation as PowerPoint writes it (ported from a deck it played). Each click is one <p:par> of the main
# sequence: the first plays clip 1, which is on screen from the start; every other one makes its clip appear and
# play at the same moment ("with previous"): the clip shows the frame where the last one stopped and goes on from
# it. No togglePause triggers: a mouse click on a video advances, like the clicker.
CLICK = ('<p:par><p:cTn id="{id}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
         '<p:childTnLst><p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
         '<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
APPEAR = ('<p:par><p:cTn id="{b}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" '
          'nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr>'
          '<p:cTn id="{c}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
          '<p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName>'
          '</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')
PLAY = ('<p:par><p:cTn id="{b}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" '
        'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        '<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{c}" dur="{dur}" fill="hold"/>'
        '<p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>')
MEDIA = ('<p:video><p:cMediaNode vol="80000"><p:cTn id="{id}" fill="hold" display="0"><p:stCondLst>'
         '<p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl>'
         '</p:cMediaNode></p:video>')
TIMING = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
          '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
          '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
          '{clicks}</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/>'
          '</p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/>'
          '</p:tgtEl></p:cond></p:nextCondLst></p:seq>{media}</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


def timing(spids, durs):
    """The <p:timing> of a slide whose clips (shape ids `spids`, durations `durs` in ms) play one per click."""
    ids = itertools.count(3)  # 1 and 2 are the root and the main sequence
    clicks = []
    for k, (sp, dur) in enumerate(zip(spids, durs)):
        cid, inner = next(ids), next(ids)
        appear = "" if k == 0 else APPEAR.format(b=next(ids), c=next(ids), sp=sp)
        play = PLAY.format(b=next(ids), c=next(ids), sp=sp, dur=dur, node="clickEffect" if k == 0 else "withEffect")
        clicks.append(CLICK.format(id=cid, inner=inner, body=appear + play))
    media = "".join(MEDIA.format(id=next(ids), sp=sp) for sp in spids)
    return TIMING.format(clicks="".join(clicks), media=media)


def build_pptx(videos, dst, bar=None, logo=None):
    """videos: [(title, clips, notes)], clips = [(mp4, poster png, duration ms, width, height)], notes = one line
    per click. One slide per video: the title on top, in white on a `bar` of that color or in INK on the white
    slide; `logo` at its right end; then the clips stacked in order under it (01 at the back)."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    for title, clips, notes in videos:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        if bar:
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, BAR_H)
            band.name = "title bar"
            band.fill.solid()
            band.fill.fore_color.rgb = RGBColor.from_string(bar)
            band.line.fill.background()
            band.shadow.inherit = False
        right = SLIDE_W - LOGO_MARGIN  # where the title must end: before the logo, if there is one
        if logo:
            pic = slide.shapes.add_picture(str(logo), 0, (BAR_H - LOGO_H) // 2, height=LOGO_H)
            pic.left = SLIDE_W - LOGO_MARGIN - pic.width
            pic.name = "logo"
            right = pic.left - LOGO_MARGIN
        box = slide.shapes.add_textbox(TITLE_X, 0, right - TITLE_X, BAR_H)
        box.name = "title"
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE  # else PowerPoint shrinks it to the text when it is edited
        box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        box.text_frame.word_wrap = True
        r = box.text_frame.paragraphs[0].add_run()
        r.text = title
        r.font.name, r.font.size, r.font.bold = "Arial", Pt(32), True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if bar else RGBColor.from_string(INK.lstrip("#"))
        spids, durs = [], []
        for k, (mp4, poster, dur, vw, vh) in enumerate(clips, 1):
            movie = slide.shapes.add_movie(str(mp4), *video_box(vw, vh), poster_frame_image=str(poster),
                                           mime_type="video/mp4")
            movie.name = f"{k:02d}"
            spids.append(movie.shape_id)
            durs.append(dur)
        old = slide._element.find(qn("p:timing"))  # add_movie writes a timing of its own: replace it
        old.addprevious(parse_xml(timing(spids, durs)))
        slide._element.remove(old)
        slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    prs.save(dst)


def check_deck(deck, videos):
    """Reopen the deck and check what PowerPoint will play, against `videos` as given to build_pptx: one slide per
    video with its title; its clips in order, under the title's space and inside the slide; the click pattern; each
    play as long as its clip; the notes. Stops with an error otherwise."""
    from lxml import etree
    from pptx import Presentation
    from pptx.oxml.ns import qn

    slides = list(Presentation(deck).slides)
    if len(slides) != len(videos):
        raise SystemExit(f"{deck.name}: {len(slides)} slides for {len(videos)} videos")
    for n, (slide, (title, clips, notes)) in enumerate(zip(slides, videos), 1):
        def bad(why):
            return SystemExit(f"{deck.name}, slide {n} ({title}): {why}")
        movies = [s for s in slide.shapes if s._element.find(".//" + qn("a:videoFile")) is not None]
        if len(movies) != len(clips):
            raise bad(f"{len(movies)} videos for {len(clips)} clips")
        for m in movies:
            if m.top < BAR_H or m.left < 0 or m.left + m.width > SLIDE_W or m.top + m.height > SLIDE_H:
                raise bad(f"video {m.name} is not under the title and inside the slide")
        xml = etree.tostring(slide._element.find(qn("p:timing"))).decode()
        kinds = re.findall(r'presetClass="(\w+)"[^>]*nodeType="(\w+)"', xml)
        want = [("mediacall", "clickEffect")] + [("entr", "clickEffect"), ("mediacall", "withEffect")] * (
            len(clips) - 1)
        if kinds != want:
            raise bad(f"click pattern {kinds}")
        plays = re.findall(r'dur="(\d+)" fill="hold"/><p:tgtEl><p:spTgt spid="(\d+)"', xml)
        if [int(sp) for _, sp in plays] != [m.shape_id for m in movies]:
            raise bad("the clips do not play in order")
        for (dur, _), (mp4, *_) in zip(plays, clips):
            frames, fps, _, _ = probe(mp4)
            if abs(int(dur) - frames / fps * 1000) > 50:
                raise bad(f"{mp4.name} plays for {dur} ms but lasts {frames / fps * 1000:.0f} ms")
        if title not in [s.text_frame.text for s in slide.shapes if s.has_text_frame]:
            raise bad("no title")
        if slide.notes_slide.notes_text_frame.text != "\n".join(notes):
            raise bad("the notes differ")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("videos", nargs="*", help="sections of pauses.toml (default: all, in file order)")
    ap.add_argument("--sheet", action="store_true", help="only the pause sheets: cut nothing")
    ap.add_argument("--gif", action="store_true", help="also each clip as a GIF that plays once")
    ap.add_argument("--pptx", action="store_true", help="also build videos/clips/deck.pptx from the clips")
    ap.add_argument("--cemep", action="store_true",
                    help=f"the CEMEP title: a #{BAR} bar, the title in white, the CEMEP symbol at its right end")
    ap.add_argument("--bar", type=bar_color, help="put the title in white on a bar of this color (default: no bar)")
    ap.add_argument("--logo", type=Path, help="a picture for the right end of the title")
    args = ap.parse_args()
    if args.cemep:
        args.bar, args.logo = args.bar or BAR, args.logo or CEMEP_LOGO
    if not PAUSES_FILE.exists():
        raise SystemExit(f"no {PAUSES_FILE}: render a video first (it writes its section there)")
    if args.logo and not args.logo.is_file():
        raise SystemExit(f"{args.logo} does not exist" + (": copy cemep_logo.png from the skill's scripts/ next to "
                                                          "cut.py" if args.logo == CEMEP_LOGO else ""))
    text = PAUSES_FILE.read_text()
    config = tomllib.loads(text)
    names = args.videos or list(config)
    for name in names:
        if name not in config:
            raise SystemExit(f"'{name}' is not in {PAUSES_FILE.name} (there: {', '.join(config)})")
        if not (OUT / f"{name}.mp4").exists():
            raise SystemExit(f"{OUT / name}.mp4 does not exist: render it first")
    deck = DEST / "deck.pptx"
    if args.pptx:
        refuse_if_open(deck)
    DEST.mkdir(parents=True, exist_ok=True)
    videos = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for name in names:
            src, pauses = OUT / f"{name}.mp4", [float(t) for t in config[name].get("pauses", [])]
            n, fps, vw, vh = probe(src)
            b = bounds(pauses, n, fps, name)
            ends = [tmp / f"{name}_end{k:02d}.png" for k in range(1, len(b))]
            for f, png in zip(b[1:], ends):
                frame_png(src, f, png)
            labels = [f"{k:02d}   stops at {t:.2f} s" for k, t in enumerate(pauses, 1)]
            labels.append(f"{len(b) - 1:02d}   to the end ({b[-1] / fps:.2f} s)")
            sheet(ends, labels, f"{name}: where each clip stops", DEST / f"{name}_pauses.png")
            if args.sheet:
                continue
            out_dir = DEST / name
            shutil.rmtree(out_dir, ignore_errors=True)  # only this video's generated clips (no stale ones)
            out_dir.mkdir()
            comments = clip_comments(text, name)
            clips, notes = [], []
            for k, (a, e) in enumerate(zip(b[:-1], b[1:]), 1):
                clip, poster = out_dir / f"{k:02d}.mp4", tmp / f"{name}_start{k:02d}.png"
                cut(src, a, e, clip)
                if args.gif:
                    to_gif(clip, clip.with_suffix(".gif"), loop=-1)
                frame_png(src, a, poster)
                clips.append((clip, poster, int(round((e - a + 1) / fps * 1000)), vw, vh))
                comment = comments[k - 1] if k <= len(comments) else ""
                notes.append(f"click {k} · {a / fps:.2f}–{e / fps:.2f} s · {comment}".rstrip(" ·"))
            videos.append((config[name].get("title", name), clips, notes))
            print(f"{name}: {len(clips)} clips in {out_dir}")
        if args.pptx and videos:
            build_pptx(videos, deck, args.bar, args.logo)
            check_deck(deck, videos)
            print(f"{deck.name}: {len(videos)} slides, {sum(len(c) for _, c, _ in videos)} clips, checked")
            print("many videos can make PowerPoint lag: on the presenting computer, File > Options > Advanced > "
                  "check 'Disable Slide Show hardware graphics acceleration'")
    print("pause sheets in", DEST)


if __name__ == "__main__":
    main()
