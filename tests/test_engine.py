"""Checks of the engine's logic (no rendering):
uv run --with pytest --with numpy --with matplotlib --with python-pptx pytest tests -q"""
import argparse
import re
import subprocess
import tomllib
import xml.etree.ElementTree as ET

import pytest

import anim
import cut


def test_new_pauses_section_has_a_title(tmp_path, monkeypatch):
    monkeypatch.setattr(anim, "PAUSES_FILE", tmp_path / "pauses.toml")
    anim.sync_pauses("intro", 12.0, [(1.9, "flow", 1), (4.5, "flow", 2)])
    section = tomllib.loads((tmp_path / "pauses.toml").read_text())["intro"]
    assert section["title"] == "intro"
    assert section["pauses"] == [1.9, 4.5]


PAUSES = '''
[intro]  # 12.0 s
title = "Intro"
pauses = [
  1.9,    # 1  flow, step 1
  4.5,    # the fitted line
          # 3  (to the end)
]

["my video"]  # 5.0 s
pauses = [
  2.0,    # 1  a, step 1
]
'''


def test_clip_comments():
    assert cut.clip_comments(PAUSES, "intro") == ["flow, step 1", "the fitted line", "(to the end)"]
    assert cut.clip_comments(PAUSES, "my video") == ["a, step 1", ""]
    assert cut.clip_comments(PAUSES, "missing") == []


EDITED = '''
# a note above the section
  [edited]  # 9.0 s
# a note above the pauses
title = "Edited"
pauses = [
  1.9,    # 1  first
  4.0,
          # 2  merged away: its number was deleted
  # 5.5,  # 3  commented out
  6.0,    # 3 standards on screen
  8.0     # 4  last pause, no comma
          # 5  (to the end)
          # a second comment line
]
'''


def test_clip_comments_follow_the_pauses_after_hand_edits():
    """Note k belongs to clip k: a split without a comment, a deleted number, a commented-out pause, no trailing
    comma, comment lines outside the list or wrapped onto a second line shift nothing."""
    assert cut.clip_comments(EDITED, "edited") == ["first", "", "3 standards on screen", "last pause, no comma",
                                                   "(to the end)"]


def test_video_box():
    assert cut.video_box(1920, 1080) == (1101795, 1239520, 9988409, 5618480)
    x, y, w, h = cut.video_box(1080, 1920)  # vertical: as tall as the room under the bar, centered
    assert y == cut.BAR_H and h == cut.SLIDE_H - cut.BAR_H and abs(2 * x + w - cut.SLIDE_W) <= 1


def test_bar_color():
    assert cut.bar_color("#232751") == "232751" and cut.bar_color("1f4e79") == "1F4E79"
    with pytest.raises(argparse.ArgumentTypeError):
        cut.bar_color("navy")


def test_timing_pattern():
    xml = cut.timing([5, 6, 7], [1000, 2000, 3000])
    ET.fromstring(xml)  # well-formed
    kinds = re.findall(r'presetClass="(\w+)"[^>]*nodeType="(\w+)"', xml)
    assert kinds == [("mediacall", "clickEffect"), ("entr", "clickEffect"), ("mediacall", "afterEffect"),
                     ("entr", "clickEffect"), ("mediacall", "afterEffect")]
    plays = re.findall(r'dur="(\d+)" fill="hold"/><p:tgtEl><p:spTgt spid="(\d+)"', xml)
    assert plays == [("1000", "5"), ("2000", "6"), ("3000", "7")]
    ids = [int(i) for i in re.findall(r'<p:cTn id="(\d+)"', xml)]
    assert ids == sorted(set(ids))  # unique, in document order
    assert "togglePause" not in xml


def make_clip(path, size="320x180"):
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", f"testsrc=size={size}:rate=30",
                    "-t", "1", "-pix_fmt", "yuv420p", "-c:v", "libx264", str(path)], check=True)
    poster = path.with_suffix(".png")
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", str(path), "-frames:v", "1", str(poster)], check=True)
    n, fps, w, h = cut.probe(path)
    return path, poster, round(n / fps * 1000), w, h


def test_deck_round_trip(tmp_path):
    a = [make_clip(tmp_path / f"a{k}.mp4") for k in (1, 2, 3)]
    b = [make_clip(tmp_path / "b1.mp4", "180x320")]  # one clip (no pauses), vertical
    videos = [("Intro", a, ["click 1 · x", "click 2", "click 3"]), ("Vertical", b, ["click 1"])]
    deck = tmp_path / "deck.pptx"
    cut.build_pptx(videos, deck, "232751", a[0][1])
    cut.check_deck(deck, videos)
    with pytest.raises(SystemExit, match="notes"):
        cut.check_deck(deck, [videos[0], ("Vertical", b, ["other"])])
    with pytest.raises(SystemExit, match="slides"):
        cut.check_deck(deck, videos[:1])


def test_title_stays_centered_in_the_bar_and_clear_of_a_wide_logo(tmp_path):
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE

    clip = make_clip(tmp_path / "a1.mp4")
    logo = tmp_path / "logo.png"  # 2:1
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=white:s=400x200",
                    "-frames:v", "1", str(logo)], check=True)
    deck = tmp_path / "deck.pptx"
    cut.build_pptx([("A long title that could run under the logo", [clip], ["click 1"])], deck, cut.BAR, logo)
    shapes = {s.name: s for s in Presentation(deck).slides[0].shapes}
    title, pic = shapes["title"], shapes["logo"]
    assert title.text_frame.auto_size == MSO_AUTO_SIZE.NONE  # PowerPoint keeps it as tall as the bar when edited
    assert title.top == 0 and title.height == cut.BAR_H
    assert title.left + title.width <= pic.left


def test_refuses_an_open_deck(tmp_path):
    deck = tmp_path / "deck.pptx"
    cut.refuse_if_open(deck)  # nothing open: fine
    (tmp_path / "~$deck.pptx").write_text("")
    with pytest.raises(SystemExit, match="open in PowerPoint"):
        cut.refuse_if_open(deck)
